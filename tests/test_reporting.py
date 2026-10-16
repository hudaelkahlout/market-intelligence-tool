from io import BytesIO
from openpyxl import load_workbook
import pytest
from market_engine import clean_reviews
from reporting import build_report

@pytest.mark.parametrize('n',[2,3,4])
def test_workbook(n):
    apps=[{'app_id':str(i),'title':f'App {i}','rating':4} for i in range(n)]
    rows=clean_reviews([{'content':'=HYPERLINK("bad")','score':1,'at':'2026-10-16'},{'content':'Very helpful','score':5}],'0')
    data={'apps':apps,'reviews':rows,'analyzed_at':'2026-10-16','persisted':True}
    book=load_workbook(BytesIO(build_report(data)))
    assert book.sheetnames==['Executive Summary','Competitor Comparison','All Reviews','Negative Reviews','AI Market Gaps']
    assert book['Competitor Comparison'].max_row==n+1
    assert book['All Reviews']['D2'].data_type=='s'
    assert book['All Reviews']['D2'].value.startswith('=HYPERLINK')
    assert book['All Reviews'].freeze_panes=='A2'
    assert book['All Reviews'].max_row==3 and book['Negative Reviews'].max_row==2
    assert book['Competitor Comparison']['G2'].number_format=='0.0%'
    assert book['All Reviews']['D2'].alignment.wrap_text
    assert book['All Reviews']['D1'].font.bold
    assert 'negative_percentage' in [c.value for c in book['Competitor Comparison'][1]]

def test_empty_and_ai():
    data={'apps':[{'app_id':'a','title':'A'},{'app_id':'b','title':'B'}], 'reviews':[], 'analyzed_at':'now',
        'ai_model':'gemini-3.5-flash-lite','ai_notice':'Requested model unavailable; fallback used',
        'ai':{'market_gaps':['Validate simpler onboarding']}}
    book=load_workbook(BytesIO(build_report(data)))
    assert book['Negative Reviews'].max_row==1
    assert book['AI Market Gaps']['B2'].value=='Validate simpler onboarding'
    executive={row[0].value:row[1].value for row in book['Executive Summary'].iter_rows(min_row=2)}
    assert executive['AI model used']=='gemini-3.5-flash-lite'
