"""Professional Excel reports generated in memory using pandas/openpyxl."""
from io import BytesIO
from pathlib import Path
import re
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from market_engine import comparison, review_frame
from market_ai import SECTIONS


def safe_text(value):
    """Remove XML control characters while preserving text as literal cells."""
    if isinstance(value,str):
        return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '',value)[:32767]
    return value

def build_report(analysis: dict, path: str | Path | None = None) -> bytes:
    """Build all five required sheets, with optional explicit file persistence."""
    apps, rows = analysis['apps'], analysis['reviews']
    summary = comparison(apps,rows)
    frame = review_frame(rows).drop(columns=['fingerprint'])
    frame['play_url'] = frame.app_id.map({a['app_id']:a.get('play_url','') for a in apps})
    if not frame.empty:
        frame['review_timestamp'] = pd.to_datetime(frame.review_timestamp,utc=True,errors='coerce').dt.tz_localize(None)
    ai = analysis.get('ai') or {}
    ai_rows = [{'Section':title,'Analysis':text} for key,title in SECTIONS.items() for text in ai.get(key,[])]
    if not ai_rows:
        ai_rows=[{'Section':'AI status','Analysis':'AI analysis has not been generated or is unavailable.'}]
    executive = [
        {'Item':'Analysis date (UTC)','Value':analysis['analyzed_at']},
        {'Item':'Selected competitors','Value':', '.join(a['title'] for a in apps)},
        {'Item':'Collected reviews','Value':len(rows)},
        {'Item':'Method','Value':'Recent US English samples; sentiment: 1–2 Negative, 3 Neutral, 4–5 Positive. Sample percentages are not market prevalence.'},
        {'Item':'Persistence','Value':'Saved to SQLite' if analysis.get('persisted') else 'Not saved'},
    ]
    if analysis.get('ai_model'):
        executive.append({'Item':'AI model used','Value':analysis['ai_model']})
    if analysis.get('ai_notice'):
        executive.append({'Item':'AI availability note','Value':analysis['ai_notice']})
    for record in summary.to_dict('records'):
        average=record['average_review_score']
        average='Unavailable' if pd.isna(average) else f'{average:.2f}'
        play='Unavailable' if pd.isna(record['play_rating']) else f"{record['play_rating']:.2f}"
        negative='Unavailable' if pd.isna(record['negative_percentage']) else f"{record['negative_percentage']:.1%}"
        executive.append({'Item':record['app_title']+' — comparison',
            'Value':f"Play rating: {play}; reviews: {record['review_count']}; sample average: {average}; negative share: {negative}"})
    executive += [{'Item':'Collection warning','Value':w} for w in analysis.get('warnings',[])]
    executive += [{'Item':r['Section'],'Value':r['Analysis']} for r in ai_rows]
    sheets={'Executive Summary':pd.DataFrame(executive), 'Competitor Comparison':summary,
        'All Reviews':frame, 'Negative Reviews':frame[frame.score.isin([1,2])],
        'AI Market Gaps':pd.DataFrame(ai_rows)}
    buffer=BytesIO()
    with pd.ExcelWriter(buffer,engine='openpyxl') as writer:
        for name,data in sheets.items():
            data=data.map(safe_text)
            data.to_excel(writer,sheet_name=name,index=False)
            ws=writer.sheets[name]
            ws.freeze_panes='A2'
            ws.auto_filter.ref=ws.dimensions
            ws.sheet_view.showGridLines=False
            for cell in ws[1]:
                cell.font=Font(bold=True,color='FFFFFF')
                cell.fill=PatternFill('solid',fgColor='17354A')
                cell.alignment=Alignment(wrap_text=True,vertical='center')
            ws.row_dimensions[1].height=32
            for col_index,column in enumerate(data.columns,1):
                width=64 if column in ('review_text','Analysis','Value') else 28 if 'timestamp' in column else max(18,min(32,len(column)+3))
                ws.column_dimensions[get_column_letter(col_index)].width=width
                for cells in ws.iter_rows(min_row=2,min_col=col_index,max_col=col_index):
                    cell=cells[0]
                    # Force all source/model strings to literal text (formula injection defense).
                    if isinstance(cell.value,str):
                        cell.data_type='s'
                    cell.alignment=Alignment(wrap_text=True,vertical='top')
                    if 'percentage' in column:
                        cell.number_format='0.0%'
                    elif column in ('play_rating','average_review_score'):
                        cell.number_format='0.00'
                    elif column=='review_timestamp':
                        cell.number_format='yyyy-mm-dd hh:mm'
            for index in range(2,ws.max_row+1):
                line_counts=[]
                for cell in ws[index]:
                    width=max(8,int(ws.column_dimensions[cell.column_letter].width)-2)
                    value=str(cell.value) if cell.value is not None else ''
                    line_counts.append(sum(max(1,(len(part)+width-1)//width) for part in value.split('\n')))
                ws.row_dimensions[index].height=min(409,max(30,15*max(line_counts)+10))
    payload=buffer.getvalue()
    if path:
        target=Path(path)
        target.parent.mkdir(parents=True,exist_ok=True)
        target.write_bytes(payload)
    return payload
